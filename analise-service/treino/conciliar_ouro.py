"""Etapa 3 do P8: concilia o conjunto OURO (rotulado à mão) em `dados_locais/ouro_final.csv`.

    cd analise-service
    python treino/conciliar_ouro.py --verificar          # só os nomes das colunas e as contagens de cada planilha
    python treino/conciliar_ouro.py                      # --modo unico (padrão por ora: João é o único rotulador)
    python treino/conciliar_ouro.py --modo pares         # 4 rotuladores em pares (1-2 = metade A, 3-4 = metade B)

`--modo unico`: lê `ouro_rotulador_1_PRONTO.xlsx` (metade A, obrigatória) e `ouro_rotulador_3_PRONTO.xlsx` (metade B;
sem ela o ouro fica "parcial"). Se existirem `ouro_reteste_PRONTO.xlsx` e `reteste_mapa.csv` (R01.. -> G..), mede o
kappa de Cohen TESTE-RETESTE (o rotulador contra ele mesmo) para `padrao` e `interesse`.
`--modo pares`: lê as 4 planilhas, mede o kappa por par e campo e grava `divergencias.xlsx` para arbitragem; com
`divergencias_PRONTO.xlsx` (coluna `decisao` preenchida) aplica a arbitragem e grava o ouro.
Kappa < 0,6 em algum campo => aviso: revisar as definições antes de usar o campo.

O ouro MEDE, nunca treina. Lê `.xlsx` com `openpyxl` (só `requirements-treino.txt`). Colunas por NOME de cabeçalho
(`--col-*`, sem distinguir maiúsculas; `--col-texto` aceita alternativas separadas por `|`). Lê SEMPRE a aba `Rotular`
(cabeçalho na linha 1: ID | contexto | contexto | TRECHO ALVO | padrao_1 | padrao_2 | interesse | duvida |
observacao) e valida `padrao_1`/`padrao_2` contra a aba `Listas` (fora da lista: só contagem). Saídas: `ouro_final.csv`, `ouro_final.json` (modo, completo/parcial, kappas) e um relatório só com métricas
em `docs/evidencias/ouro-<data>.md`. Nunca imprime texto de trecho.
"""
import argparse
import datetime as dt
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from treino.comum import (  # noqa: E402
    DADOS, EVIDENCIAS, SEM_PADRAO, gravar_csv, grupos_dos_padroes, imprimir, ler_csv, normalizar_padrao,
    normalizar_sim_nao, sem_acento,
)

LIMIAR_KAPPA = 0.6
CAMPOS_KAPPA = ("padrao", "interesse")
CAMPOS_DIVERGENCIA = ("padrao_1", "padrao_2", "interesse")
COLUNAS_OURO = ["id", "texto", "padrao_1", "padrao_2", "risco", "oportunidade", "interesse", "duvida", "metade"]
ARQUIVOS = {
    1: "ouro_rotulador_1_PRONTO.xlsx", 2: "ouro_rotulador_2_PRONTO.xlsx",
    3: "ouro_rotulador_3_PRONTO.xlsx", 4: "ouro_rotulador_4_PRONTO.xlsx",
}
RETESTE, MAPA_RETESTE = "ouro_reteste_PRONTO.xlsx", "reteste_mapa.csv"
DIVERGENCIAS, DIVERGENCIAS_PRONTO = "divergencias.xlsx", "divergencias_PRONTO.xlsx"
ABA_ROTULAR, ABA_LISTAS = "Rotular", "Listas"
COL_ID, COL_TEXTO = "ID", "TRECHO ALVO|texto"  # layout real da planilha | layout antigo
LIMITACAO_UNICO = ("Rotulador único (João Paulo), sem validação cruzada entre pessoas; ele também escreveu as "
                   "expressões regulares. A concordância medida é teste-reteste (consigo mesmo).")


class Colunas:
    def __init__(self, args: argparse.Namespace | None = None):
        a = args or argparse.Namespace(col_id=COL_ID, col_texto=COL_TEXTO, col_padrao1="padrao_1",
                                       col_padrao2="padrao_2", col_interesse="interesse", col_duvida="duvida")
        self.id, self.padrao1 = a.col_id, a.col_padrao1
        self.texto = [_chave(n) for n in a.col_texto.split("|") if n.strip()]
        self.padrao2, self.interesse, self.duvida = a.col_padrao2, a.col_interesse, a.col_duvida


# ------------------------------------------------------------------ leitura

def _chave(nome: object) -> str:
    return sem_acento(str(nome or "")).replace(" ", "_")


def _slug_padrao(valor: object) -> str:
    return sem_acento(str(valor or "")).replace(" ", "_").replace("-", "_")


def ler_planilha(caminho: Path, col_id: str, aba: str | None = ABA_ROTULAR) -> list[dict[str, object]]:
    """Linhas da aba `aba` (None = a primeira) como dicts pelo cabeçalho (achado nas 10 primeiras linhas pela coluna
    de id). Colunas ocultas (a C do layout real) são lidas normalmente."""
    import openpyxl

    livro = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        if aba is None:
            planilha = livro.worksheets[0]
        elif aba in livro.sheetnames:
            planilha = livro[aba]
        else:
            raise SystemExit(f"erro: {caminho.name} sem a aba {aba!r} (abas: {', '.join(livro.sheetnames)})")
        linhas = list(planilha.iter_rows(values_only=True))
    finally:
        livro.close()
    for i, linha in enumerate(linhas[:10]):
        nomes = [_chave(c) for c in linha]
        if _chave(col_id) in nomes:
            dados = []
            for valores in linhas[i + 1:]:
                registro = {nome: valor for nome, valor in zip(nomes, valores) if nome}
                if str(registro.get(_chave(col_id)) or "").strip():
                    dados.append(registro)
            return dados
    raise SystemExit(f"erro: {caminho.name} sem a coluna {col_id!r} nas 10 primeiras linhas (use --verificar)")


def lista_de_padroes(caminho: Path) -> set[str] | None:
    """Vocabulário permitido para padrao_1/padrao_2: todo valor textual da aba `Listas`, normalizado.
    None se a planilha não tem a aba (então não há o que validar)."""
    import openpyxl

    livro = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        if ABA_LISTAS not in livro.sheetnames:
            return None
        return {_slug_padrao(v) for linha in livro[ABA_LISTAS].iter_rows(values_only=True) for v in linha
                if isinstance(v, str) and v.strip()}
    finally:
        livro.close()


def cabecalho_planilha(caminho: Path, col_id: str) -> dict:
    linhas = ler_planilha(caminho, col_id)
    lista = lista_de_padroes(caminho)
    return {"colunas": sorted(linhas[0]) if linhas else [], "linhas": len(linhas),
            "valores_na_aba_listas": None if lista is None else len(lista)}


def _padrao(valor: object, permitidos: set[str] | None, invalidos: dict) -> str | None:
    """Padrão normalizado. Fora da aba Listas ou não-padrão => None, só contado (o valor nunca é impresso)."""
    padrao = normalizar_padrao(valor)
    if padrao == SEM_PADRAO:
        return padrao
    if permitidos is not None and _slug_padrao(valor) not in permitidos:
        invalidos["fora_da_lista"] += 1
        return None
    invalidos["padrao"] += padrao is None
    return padrao


def _texto(linha: dict, c: Colunas) -> str:
    return next((str(linha[n]) for n in c.texto if linha.get(n) not in (None, "")), "")


def rotulos(linhas: list[dict], c: Colunas, metade: str,
            permitidos: set[str] | None = None) -> tuple[dict[str, dict], dict]:
    """id -> rótulo normalizado; e contagens de valores inválidos (sem imprimir os valores)."""
    saida: dict[str, dict] = {}
    invalidos = {"padrao": 0, "interesse": 0, "ids_repetidos": 0, "fora_da_lista": 0}
    for linha in linhas:
        id_ = str(linha.get(_chave(c.id))).strip()
        if id_ in saida:
            invalidos["ids_repetidos"] += 1
            continue
        p1 = _padrao(linha.get(_chave(c.padrao1)), permitidos, invalidos)
        p2 = _padrao(linha.get(_chave(c.padrao2)), permitidos, invalidos)
        interesse = normalizar_sim_nao(linha.get(_chave(c.interesse)))
        invalidos["interesse"] += interesse is None
        saida[id_] = {
            "id": id_, "texto": _texto(linha, c), "metade": metade,
            "padrao_1": p1 or "", "padrao_2": p2 or "", "interesse": interesse,
            "duvida": bool(normalizar_sim_nao(linha.get(_chave(c.duvida)))),
        }
    return saida, invalidos


# ------------------------------------------------------------------ kappa

def kappa_cohen(a: list, b: list) -> float | None:
    """Kappa de Cohen para dois rótulos categóricos pareados. None sem pares."""
    pares = [(x, y) for x, y in zip(a, b) if x not in (None, "") and y not in (None, "")]
    if not pares:
        return None
    n = len(pares)
    po = sum(1 for x, y in pares if x == y) / n
    categorias = {x for x, _ in pares} | {y for _, y in pares}
    pe = sum((sum(1 for x, _ in pares if x == k) / n) * (sum(1 for _, y in pares if y == k) / n) for k in categorias)
    if pe >= 1:
        return 1.0 if po == 1 else 0.0
    return round((po - pe) / (1 - pe), 4)


def _valor_kappa(rotulo: dict, campo: str):
    if campo == "padrao":
        return rotulo["padrao_1"] or None
    return rotulo["interesse"]


def kappas(r1: dict[str, dict], r2: dict[str, dict], ids: list[str]) -> dict[str, float | None]:
    return {campo: kappa_cohen([_valor_kappa(r1[i], campo) for i in ids], [_valor_kappa(r2[i], campo) for i in ids])
            for campo in CAMPOS_KAPPA}


def avisos_de_kappa(medidos: dict[str, dict[str, float | None]]) -> list[str]:
    avisos = []
    for origem, por_campo in medidos.items():
        for campo, valor in por_campo.items():
            if valor is not None and valor < LIMIAR_KAPPA:
                avisos.append(f"kappa {origem}/{campo} = {valor} < {LIMIAR_KAPPA}: revisar as definições de "
                              f"'{campo}' antes de usar este campo")
    return avisos


# ------------------------------------------------------------------ saída

def linha_ouro(r: dict) -> dict:
    padroes = [p for p in (r["padrao_1"], r["padrao_2"]) if p and p != SEM_PADRAO]
    grupos = grupos_dos_padroes(padroes)
    return {**r, "risco": "1" if grupos["risco"] else "0", "oportunidade": "1" if grupos["oportunidade"] else "0",
            "interesse": "" if r["interesse"] is None else ("1" if r["interesse"] else "0"),
            "duvida": "1" if r["duvida"] else "0"}


def gravar_ouro(pasta: Path, ouro: list[dict], meta: dict) -> None:
    gravar_csv(pasta / "ouro_final.csv", [linha_ouro(r) for r in ouro], COLUNAS_OURO)
    (pasta / "ouro_final.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")


def resumo_do_ouro(ouro: list[dict]) -> dict:
    linhas = [linha_ouro(r) for r in ouro]
    return {"trechos": len(linhas), "duvida": sum(1 for x in linhas if x["duvida"] == "1"),
            "risco": sum(1 for x in linhas if x["risco"] == "1"),
            "oportunidade": sum(1 for x in linhas if x["oportunidade"] == "1"),
            "interesse": sum(1 for x in linhas if x["interesse"] == "1"),
            "por_padrao": {p: sum(1 for r in ouro if p in (r["padrao_1"], r["padrao_2"]))
                           for p in sorted({r["padrao_1"] for r in ouro} | {r["padrao_2"] for r in ouro}) if p}}


def escrever_relatorio(caminho: Path, meta: dict) -> None:
    linhas = [
        f"# Conciliação do conjunto ouro — {meta['data']}", "",
        "Só métricas e contagens; nenhum texto de trecho. Gerado por `treino/conciliar_ouro.py`.", "",
        f"- Modo: **{meta['modo']}**" + (f" — {LIMITACAO_UNICO}" if meta["modo"] == "unico" else ""),
        f"- Ouro: **{'completo' if meta['completo'] else 'PARCIAL'}** (metades: {', '.join(meta['metades'])})",
        f"- Trechos: {meta['resumo']['trechos']} (com dúvida: {meta['resumo']['duvida']}); positivos: risco "
        f"{meta['resumo']['risco']}, oportunidade {meta['resumo']['oportunidade']}, interesse "
        f"{meta['resumo']['interesse']}",
        f"- Valores inválidos (contagem): {json.dumps(meta['invalidos'], ensure_ascii=False)}", "",
        "## Concordância (kappa de Cohen)", "", "| Origem | padrao | interesse | pares |", "|---|---|---|---|",
    ]
    for origem, k in meta["kappas"].items():
        linhas.append(f"| {origem} | {k.get('padrao')} | {k.get('interesse')} | {meta['pares_kappa'].get(origem, 0)} |")
    if not meta["kappas"]:
        linhas.append("| (não medido) | — | — | 0 |")
    linhas += ["", "## Avisos", ""] + ([f"- {a}" for a in meta["avisos"]] or ["- nenhum"])
    linhas += ["", "## Trechos por padrão no ouro (indicativo, ~8 por padrão)", "", "| Padrão | trechos |", "|---|---|"]
    linhas += [f"| {p} | {n} |" for p, n in meta["resumo"]["por_padrao"].items()]
    caminho.parent.mkdir(parents=True, exist_ok=True)
    caminho.write_text("\n".join(linhas) + "\n", encoding="utf-8")


# ------------------------------------------------------------------ modos

def _somar(invalidos: dict, novo: dict) -> None:
    for k, v in novo.items():
        invalidos[k] = invalidos.get(k, 0) + v


def _rotulos_do_arquivo(caminho: Path, c: Colunas, metade: str) -> tuple[dict[str, dict], dict]:
    return rotulos(ler_planilha(caminho, c.id), c, metade, lista_de_padroes(caminho))


def modo_unico(pasta: Path, c: Colunas) -> tuple[list[dict] | None, dict]:
    if not (pasta / ARQUIVOS[1]).is_file():
        raise SystemExit(f"erro: falta {ARQUIVOS[1]} (metade A) em {pasta.name}/")
    invalidos: dict = {}
    a, inv = _rotulos_do_arquivo(pasta / ARQUIVOS[1], c, "A")
    _somar(invalidos, inv)
    ouro, metades = dict(a), ["A"]
    if (pasta / ARQUIVOS[3]).is_file():
        b, inv = _rotulos_do_arquivo(pasta / ARQUIVOS[3], c, "B")
        _somar(invalidos, inv)
        if set(a) & set(b):
            raise SystemExit("erro: há ids repetidos entre a metade A e a metade B")
        ouro.update(b)
        metades.append("B")
    medidos, pares = {}, {}
    if (pasta / RETESTE).is_file() and (pasta / MAPA_RETESTE).is_file():
        reteste, inv = _rotulos_do_arquivo(pasta / RETESTE, c, "reteste")
        mapa = _ler_mapa(pasta / MAPA_RETESTE)
        ids_r = [r for r in reteste if mapa.get(r) in ouro]
        if ids_r:
            original = {r: ouro[mapa[r]] for r in ids_r}
            medidos["teste-reteste"] = kappas(original, reteste, ids_r)
            pares["teste-reteste"] = len(ids_r)
    meta = {"modo": "unico", "completo": metades == ["A", "B"], "metades": metades, "kappas": medidos,
            "pares_kappa": pares, "invalidos": invalidos, "avisos": avisos_de_kappa(medidos),
            "limitacao": LIMITACAO_UNICO}
    if not medidos:
        meta["avisos"].append("teste-reteste ainda não medido (faltam ouro_reteste_PRONTO.xlsx e reteste_mapa.csv)")
    return list(ouro.values()), meta


def _ler_mapa(caminho: Path) -> dict[str, str]:
    """reteste_mapa.csv: colunas `reteste` e `ouro` (ou as duas primeiras), R01 -> G...."""
    linhas = ler_csv(caminho)
    if not linhas:
        return {}
    nomes = list(linhas[0])
    col_r = "reteste" if "reteste" in nomes else nomes[0]
    col_o = "ouro" if "ouro" in nomes else nomes[1]
    return {str(x[col_r]).strip(): str(x[col_o]).strip() for x in linhas}


def _divergencias(r1: dict[str, dict], r2: dict[str, dict]) -> list[dict]:
    saida = []
    for id_ in r1:
        for campo in CAMPOS_DIVERGENCIA:
            v1, v2 = r1[id_][campo], r2[id_][campo]
            if v1 != v2:
                saida.append({"id": id_, "texto": r1[id_]["texto"], "campo": campo,
                              "rotulador_a": _para_planilha(v1), "rotulador_b": _para_planilha(v2), "decisao": ""})
    return saida


def _para_planilha(valor) -> str:
    if isinstance(valor, bool):
        return "sim" if valor else "nao"
    return "" if valor is None else str(valor)


def _gravar_divergencias(caminho: Path, divergencias: list[dict]) -> None:
    import openpyxl

    livro = openpyxl.Workbook()
    aba = livro.active
    colunas = ["id", "texto", "campo", "rotulador_a", "rotulador_b", "decisao"]
    aba.append(colunas)
    for d in divergencias:
        aba.append([d[k] for k in colunas])
    livro.save(caminho)


def _aplicar_decisao(rotulo: dict, campo: str, decisao: object) -> bool:
    valor = normalizar_sim_nao(decisao) if campo == "interesse" else normalizar_padrao(decisao)
    if valor is None:
        return False
    rotulo[campo] = valor
    return True


def modo_pares(pasta: Path, c: Colunas) -> tuple[list[dict] | None, dict]:
    presentes = {k for k, nome in ARQUIVOS.items() if (pasta / nome).is_file()}
    if not {1, 2} <= presentes:
        raise SystemExit(f"erro: o modo pares precisa de {ARQUIVOS[1]} e {ARQUIVOS[2]} (metade A)")
    invalidos: dict = {}
    lidos = {}
    for k in sorted(presentes):
        lidos[k], inv = _rotulos_do_arquivo(pasta / ARQUIVOS[k], c, "A" if k in (1, 2) else "B")
        _somar(invalidos, inv)
    pares = [(1, 2)] + ([(3, 4)] if {3, 4} <= presentes else [])
    medidos, n_pares, divergencias, ouro = {}, {}, [], {}
    for x, y in pares:
        if set(lidos[x]) != set(lidos[y]):
            raise SystemExit(f"erro: rotuladores {x} e {y} não têm os mesmos ids")
        ids = sorted(lidos[x])
        medidos[f"rotuladores {x}x{y}"] = kappas(lidos[x], lidos[y], ids)
        n_pares[f"rotuladores {x}x{y}"] = len(ids)
        divergencias += _divergencias(lidos[x], lidos[y])
        for id_ in ids:  # base: rotulador da esquerda; dúvida de qualquer um dos dois conta
            ouro[id_] = {**lidos[x][id_], "duvida": lidos[x][id_]["duvida"] or lidos[y][id_]["duvida"]}
    metades = ["A"] + (["B"] if (3, 4) in pares else [])
    meta = {"modo": "pares", "completo": metades == ["A", "B"], "metades": metades, "kappas": medidos,
            "pares_kappa": n_pares, "invalidos": invalidos, "avisos": avisos_de_kappa(medidos),
            "divergencias": len(divergencias)}
    if divergencias:
        pronto = pasta / DIVERGENCIAS_PRONTO
        if not pronto.is_file():
            _gravar_divergencias(pasta / DIVERGENCIAS, divergencias)
            meta["avisos"].append(f"{len(divergencias)} divergência(s) gravadas em {DIVERGENCIAS}: arbitre (coluna "
                                  f"decisao) e salve como {DIVERGENCIAS_PRONTO}")
            return None, meta
        decisoes = {(str(d.get("id")).strip(), _chave(d.get("campo"))): d.get("decisao")
                    for d in ler_planilha(pronto, "id", aba=None)}
        faltando = 0
        for d in divergencias:
            decisao = decisoes.get((d["id"], d["campo"]))
            if decisao is None or str(decisao).strip() == "" or not _aplicar_decisao(ouro[d["id"]], d["campo"], decisao):
                faltando += 1
        if faltando:
            raise SystemExit(f"erro: {faltando} divergência(s) sem decisão válida em {DIVERGENCIAS_PRONTO}")
    return list(ouro.values()), meta


def main(argv: list[str]) -> int:
    p = argparse.ArgumentParser(description="Etapa 3 do P8: concilia o conjunto ouro (só métricas na saída).")
    p.add_argument("--modo", choices=["unico", "pares"], default="unico")
    p.add_argument("--pasta", type=Path, default=DADOS, help="onde estão as planilhas (padrão dados_locais/)")
    p.add_argument("--relatorio", type=Path, default=None,
                   help="relatório só com métricas (padrão docs/evidencias/ouro-<data>.md)")
    p.add_argument("--verificar", action="store_true", help="imprime só nomes de colunas e contagens das planilhas")
    p.add_argument("--col-id", default=COL_ID, help="coluna do id (sem distinguir maiúsculas)")
    p.add_argument("--col-texto", default=COL_TEXTO,
                   help="coluna com o texto do turno ALVO; alternativas separadas por | (a primeira preenchida vale)")
    p.add_argument("--col-padrao1", default="padrao_1")
    p.add_argument("--col-padrao2", default="padrao_2")
    p.add_argument("--col-interesse", default="interesse")
    p.add_argument("--col-duvida", default="duvida")
    args = p.parse_args(argv)
    c = Colunas(args)

    if args.verificar:
        resumo = {}
        for nome in [*ARQUIVOS.values(), RETESTE, DIVERGENCIAS_PRONTO]:
            if (args.pasta / nome).is_file():
                resumo[nome] = cabecalho_planilha(args.pasta / nome, c.id)
        imprimir(resumo or {"aviso": "nenhuma planilha *_PRONTO.xlsx encontrada"})
        return 0

    ouro, meta = (modo_unico if args.modo == "unico" else modo_pares)(args.pasta, c)
    meta["data"] = dt.date.today().isoformat()
    if ouro is None:
        imprimir({k: v for k, v in meta.items() if k != "limitacao"})
        return 3
    meta["resumo"] = resumo_do_ouro(ouro)
    gravar_ouro(args.pasta, ouro, meta)
    relatorio = args.relatorio or EVIDENCIAS / f"ouro-{meta['data']}.md"
    escrever_relatorio(relatorio, meta)
    imprimir({"saida": "ouro_final.csv", "relatorio": relatorio.name, **{k: v for k, v in meta.items()
                                                                      if k not in ("limitacao",)}})
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
