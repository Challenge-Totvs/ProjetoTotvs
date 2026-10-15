"""Comparação cega entre o modelo local atual e um candidato, revisando só onde os dois discordam.

Para quando o ouro é pequeno demais para medir F1 com confiança: onde atual e candidato concordam, trocar não muda
nada; toda a diferença está nas divergências. Uma pessoa revisa uma amostra delas SEM ver o que cada modelo disse
(só responde "é risco? sim/não/dúvida"), e o placar diz qual modelo acertou mais.

    cd analise-service
    python treino/comparar_modelos.py montar                      # candidato mais recente
    python treino/comparar_modelos.py montar --candidato dados_locais/candidatos/2026-10-14 --amostra 50
    # preencher a coluna "resposta" de dados_locais/comparacao/<candidato>/revisao.xlsx
    python treino/comparar_modelos.py placar                      # só o resultado
    python treino/comparar_modelos.py placar --promover           # promove o grupo em que o candidato venceu
    python treino/comparar_modelos.py placar --promover --forcar  # promove mesmo sem vitória clara

Trechos usados: reuniões do HOLDOUT que NÃO estão no ouro (`no_ouro == 0`). O candidato nunca treinou com elas e o
ouro (que mede, não calibra) fica de fora. Grupos:
- risco: classificador atual × candidato de risco, os dois com `LIMIAR_RISCO`;
- interesse: filtro atual (palavra-chave E risco do atual < `LIMIAR_INTERESSE`) × candidato com `LIMIAR_P_INTERESSE`.

Saídas com texto ficam só em `dados_locais/comparacao/` (fora do git): a planilha de revisão e a chave. O terminal e o
relatório em docs/evidencias/ têm só contagens e o placar. Nunca imprime trecho.
"""
import argparse
import datetime as dt
import json
import math
import random
import shutil
import sys
from pathlib import Path
from typing import Any, Sequence

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.engines.modelo_local_engine import LIMIAR_INTERESSE, LIMIAR_P_INTERESSE, LIMIAR_RISCO  # noqa: E402
from app.engines.regras_engine import EXPRESSAO_INTERESSE, dobrar  # noqa: E402
from treino.comum import (  # noqa: E402
    CAMINHO_MODELO_ATUAL, DADOS, EVIDENCIAS, RAIZ, dividir_turnos, imprimir, ler_corpus, ler_csv,
)

GRUPOS = ("risco", "interesse")
PERGUNTA = {"risco": "Este trecho mostra um RISCO para a venda ou para a conta?",
            "interesse": "Este trecho mostra INTERESSE do cliente?"}
RESPOSTAS_VALIDAS = {"sim": 1, "s": 1, "nao": 0, "não": 0, "n": 0}
RESPOSTAS_DUVIDA = {"duvida", "dúvida", "d", "?"}
TAMANHO_MINIMO = 25        # trecho curto ("ok", "perfeito") não ajuda a decidir
CONTEXTO = 2               # turnos anteriores mostrados na revisão
AMOSTRA_PADRAO = 40        # divergências por grupo
SEMENTE = 20261015
ALFA = 0.05                # teste do sinal, bilateral
DESTINO = {"risco": RAIZ / "app" / "modelos" / "classificador_risco.joblib",
           "interesse": RAIZ / "app" / "modelos" / "classificador_interesse.joblib"}
COLUNAS_PLANILHA = ["id", "grupo", "pergunta", "contexto", "trecho", "resposta"]


# ------------------------------------------------------------------ trechos

def trechos_do_holdout(corpus: list[dict], holdout: list[dict], col_reuniao: str, col_texto: str,
                       col_holdout: str) -> tuple[list[dict], dict]:
    """Turnos das reuniões do holdout fora do ouro. Devolve (trechos, contagens)."""
    alvo = {(h.get(col_holdout) or "").strip() for h in holdout
            if str(h.get("no_ouro", "")).strip() in ("0", "false", "False", "")} - {""}
    contagem = {"reunioes_holdout_fora_do_ouro": len(alvo), "reunioes_encontradas": 0, "trechos": 0,
                "trechos_curtos_ignorados": 0}
    trechos, vistos = [], set()
    for linha in corpus:
        reuniao = (linha.get(col_reuniao) or "").strip()
        if reuniao not in alvo or reuniao in vistos:
            continue
        vistos.add(reuniao)
        contagem["reunioes_encontradas"] += 1
        turnos = dividir_turnos(linha.get(col_texto) or "")
        for i, (_, texto) in enumerate(turnos):
            if len(texto.strip()) < TAMANHO_MINIMO:
                contagem["trechos_curtos_ignorados"] += 1
                continue
            contexto = " || ".join(t for _, t in turnos[max(0, i - CONTEXTO):i])
            trechos.append({"id": f"{reuniao}:{i + 1}", "reuniao": reuniao, "texto": texto, "contexto": contexto})
    contagem["trechos"] = len(trechos)
    return trechos, contagem


# ------------------------------------------------------------------ divergências

def decisoes(textos: Sequence[str], atual: Any, cand_risco: Any, cand_interesse: Any | None) -> dict:
    """Por grupo, a decisão 0/1 de cada lado para cada texto."""
    p_atual = atual.predict_proba(list(textos))[:, 1]
    p_risco = cand_risco.predict_proba(list(textos))[:, 1]
    saida = {"risco": {"atual": [int(p >= LIMIAR_RISCO) for p in p_atual],
                       "candidato": [int(p >= LIMIAR_RISCO) for p in p_risco]}}
    if cand_interesse is not None:
        p_int = cand_interesse.predict_proba(list(textos))[:, 1]
        saida["interesse"] = {
            "atual": [int(bool(EXPRESSAO_INTERESSE.search(dobrar(t))) and p < LIMIAR_INTERESSE)
                      for t, p in zip(textos, p_atual)],
            "candidato": [int(p >= LIMIAR_P_INTERESSE) for p in p_int]}
    return saida


def divergencias(dec: dict) -> dict[str, list[int]]:
    """Índices em que atual e candidato discordam, por grupo."""
    return {g: [i for i, (a, c) in enumerate(zip(d["atual"], d["candidato"])) if a != c] for g, d in dec.items()}


def amostrar(indices: list[int], n: int, semente: int = SEMENTE) -> list[int]:
    """Amostra aleatória simples (reprodutível) de até n índices, em ordem embaralhada."""
    rng = random.Random(semente)
    escolhidos = indices if len(indices) <= n else rng.sample(indices, n)
    escolhidos = list(escolhidos)
    rng.shuffle(escolhidos)
    return escolhidos


# ------------------------------------------------------------------ placar

def teste_do_sinal(vitorias_a: int, vitorias_b: int) -> float:
    """p-valor bilateral exato do teste do sinal (empates/dúvidas já fora)."""
    n = vitorias_a + vitorias_b
    if n == 0:
        return 1.0
    k = min(vitorias_a, vitorias_b)
    cauda = sum(math.comb(n, i) for i in range(k + 1)) / 2 ** n
    return min(1.0, 2 * cauda)


def ler_resposta(valor: object) -> int | None:
    """1 = sim, 0 = não, None = dúvida ou em branco."""
    texto = str(valor or "").strip().lower()
    if texto in RESPOSTAS_VALIDAS:
        return RESPOSTAS_VALIDAS[texto]
    return None


def placar(respostas: list[dict], chave: dict) -> dict:
    """respostas: [{id, grupo, resposta}]; chave: {grupo: {id: {"atual": 0/1, "candidato": 0/1}}}."""
    resultado = {}
    for grupo in GRUPOS:
        if grupo not in chave:
            continue
        linhas = [r for r in respostas if r.get("grupo") == grupo]
        c = a = duvida = branco = 0
        por_direcao = {"so_candidato_marca": {"candidato": 0, "atual": 0},
                       "so_atual_marca": {"candidato": 0, "atual": 0}}
        for r in linhas:
            lados = chave[grupo].get(str(r["id"]))
            if lados is None:
                continue
            bruto = str(r.get("resposta") or "").strip().lower()
            if not bruto:
                branco += 1
                continue
            verdade = ler_resposta(bruto)
            if verdade is None:
                duvida += 1
                continue
            direcao = "so_candidato_marca" if lados["candidato"] == 1 else "so_atual_marca"
            if lados["candidato"] == verdade:
                c += 1
                por_direcao[direcao]["candidato"] += 1
            else:
                a += 1
                por_direcao[direcao]["atual"] += 1
        p = teste_do_sinal(c, a)
        if c + a == 0:
            veredito = "sem respostas"
        elif c > a and p < ALFA:
            veredito = "candidato melhor"
        elif a > c and p < ALFA:
            veredito = "atual melhor"
        else:
            veredito = "inconclusivo"
        resultado[grupo] = {"revisados": c + a, "duvida": duvida, "em_branco": branco,
                            "vitorias_candidato": c, "vitorias_atual": a,
                            "taxa_candidato": round(c / (c + a), 4) if c + a else None,
                            "p_valor_sinal": round(p, 4), "veredito": veredito, "por_direcao": por_direcao,
                            "divergencias_totais": chave.get("_totais", {}).get(grupo)}
    return resultado


# ------------------------------------------------------------------ arquivos

def _pasta_candidato(arg: Path | None, candidatos: Path) -> Path:
    if arg is not None:
        return arg
    pastas = sorted(p for p in candidatos.iterdir() if p.is_dir()) if candidatos.is_dir() else []
    if not pastas:
        raise SystemExit("erro: nenhum candidato em dados_locais/candidatos/ (rode treino/treinar.py)")
    return pastas[-1]


def gravar_planilha(caminho: Path, linhas: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "revisao"
    ws.append(COLUNAS_PLANILHA)
    for c in ws[1]:
        c.font = Font(bold=True)
    for linha in linhas:
        ws.append([linha[k] for k in COLUNAS_PLANILHA])
    larguras = {"A": 22, "B": 10, "C": 34, "D": 60, "E": 70, "F": 12}
    for col, w in larguras.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"sim,nao,duvida"', allow_blank=True)
    ws.add_data_validation(dv)
    if linhas:
        dv.add(f"F2:F{len(linhas) + 1}")
    ws.freeze_panes = "A2"
    instr = wb.create_sheet("como_preencher")
    for texto in [
        "Responda só a pergunta da linha, olhando o trecho (o contexto são as falas anteriores).",
        "Coluna 'resposta': sim, nao ou duvida. Em branco = não revisado.",
        "Você não vê o que cada modelo disse, de propósito: a revisão é cega.",
        "Depois: python treino/comparar_modelos.py placar",
    ]:
        instr.append([texto])
    instr.column_dimensions["A"].width = 100
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


def ler_planilha(caminho: Path) -> list[dict]:
    from openpyxl import load_workbook

    ws = load_workbook(caminho, read_only=True)["revisao"]
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return []
    cab = [str(c) for c in linhas[0]]
    return [dict(zip(cab, l)) for l in linhas[1:] if l and l[0] is not None]


def escrever_relatorio(caminho: Path, candidato: str, resultado: dict, contagens: dict) -> None:
    hoje = dt.date.today().isoformat()
    linhas = [f"# Comparação cega atual × candidato `{candidato}` ({hoje})", "",
              "Revisão humana só dos trechos em que os dois modelos discordam, sem saber o que cada um disse. "
              "Trechos de reuniões do holdout fora do ouro. Só contagens; nenhum texto de trecho.", "",
              f"- Trechos avaliados pelos dois modelos: {contagens.get('trechos', '—')}",
              f"- Teste do sinal bilateral, alfa = {ALFA}", "",
              "| Grupo | Divergências totais | Revisados | Candidato acertou | Atual acertou | Dúvida | p | Veredito |",
              "| --- | --- | --- | --- | --- | --- | --- | --- |"]
    for g, r in resultado.items():
        linhas.append(f"| {g} | {r['divergencias_totais'] or '—'} | {r['revisados']} | {r['vitorias_candidato']} | "
                      f"{r['vitorias_atual']} | {r['duvida']} | {r['p_valor_sinal']} | {r['veredito']} |")
    linhas += ["", "Por direção (quem acertou quando só um dos lados marcou):", ""]
    for g, r in resultado.items():
        d = r["por_direcao"]
        linhas.append(f"- {g}: só o candidato marcou → candidato {d['so_candidato_marca']['candidato']} × "
                      f"atual {d['so_candidato_marca']['atual']}; só o atual marcou → candidato "
                      f"{d['so_atual_marca']['candidato']} × atual {d['so_atual_marca']['atual']}")
    caminho.parent.mkdir(parents=True, exist_ok=True)
    caminho.write_text("\n".join(linhas) + "\n", encoding="utf-8")


def promover(grupos: list[str], pasta: Path, backup: Path) -> list[str]:
    """Copia o candidato por cima do modelo em uso, guardando o anterior em `backup`. Devolve o que foi feito."""
    feito = []
    backup.mkdir(parents=True, exist_ok=True)
    for g in grupos:
        origem, destino = pasta / f"{g}.joblib", DESTINO[g]
        if destino.is_file():
            shutil.copy2(destino, backup / destino.name)
            feito.append(f"{g}: anterior guardado em {backup / destino.name}")
        shutil.copy2(origem, destino)
        feito.append(f"{g}: {origem} -> {destino}")
    return feito


# ------------------------------------------------------------------ comandos

def cmd_montar(args) -> int:
    import joblib

    pasta = _pasta_candidato(args.candidato, args.candidatos)
    trechos, contagem = trechos_do_holdout(ler_corpus(args.corpus), ler_csv(args.holdout), args.col_reuniao,
                                           args.col_texto, args.col_holdout)
    if not trechos:
        raise SystemExit("erro: nenhum trecho do holdout fora do ouro (confira --col-reuniao/--col-holdout)")
    # joblib é pickle: só o artefato do repo e candidatos gerados por treino/treinar.py
    atual = joblib.load(args.atual)
    cand_risco = joblib.load(pasta / "risco.joblib")
    cand_int = joblib.load(pasta / "interesse.joblib") if "interesse" in args.grupos else None
    textos = [t["texto"] for t in trechos]
    dec = decisoes(textos, atual, cand_risco, cand_int)
    dec = {g: d for g, d in dec.items() if g in args.grupos}
    div = divergencias(dec)

    linhas, chave = [], {"_candidato": pasta.name, "_totais": {g: len(v) for g, v in div.items()}}
    resumo_div = {}
    for g in args.grupos:
        idx = div[g]
        so_cand = sum(1 for i in idx if dec[g]["candidato"][i] == 1)
        resumo_div[g] = {"divergencias": len(idx), "so_candidato_marca": so_cand,
                         "so_atual_marca": len(idx) - so_cand,
                         "atual_marca": sum(dec[g]["atual"]), "candidato_marca": sum(dec[g]["candidato"])}
        chave[g] = {}
        for i in amostrar(idx, args.amostra, args.semente):
            t = trechos[i]
            rid = f"{g[0].upper()}-{t['id']}"
            linhas.append({"id": rid, "grupo": g, "pergunta": PERGUNTA[g], "contexto": t["contexto"],
                           "trecho": t["texto"], "resposta": ""})
            chave[g][rid] = {"atual": dec[g]["atual"][i], "candidato": dec[g]["candidato"][i]}
        resumo_div[g]["na_amostra"] = len(chave[g])
    random.Random(args.semente).shuffle(linhas)  # mistura os grupos na planilha

    saida = args.saida or (DADOS / "comparacao" / pasta.name)
    gravar_planilha(saida / "revisao.xlsx", linhas)
    (saida / "chave.json").write_text(json.dumps(chave, ensure_ascii=False, indent=1), encoding="utf-8")
    (saida / "contagens.json").write_text(json.dumps({**contagem, "grupos": resumo_div}, indent=1), encoding="utf-8")
    imprimir({"candidato": pasta.name, "trechos": contagem, "grupos": resumo_div,
              "planilha": str((saida / "revisao.xlsx").relative_to(RAIZ)) if saida.is_relative_to(RAIZ) else str(saida)})
    print("preencha a coluna 'resposta' (sim / nao / duvida) e rode: python treino/comparar_modelos.py placar")
    return 0


def cmd_placar(args) -> int:
    pasta = _pasta_candidato(args.candidato, args.candidatos)
    saida = args.saida or (DADOS / "comparacao" / pasta.name)
    chave_arq, planilha = saida / "chave.json", saida / "revisao.xlsx"
    if not chave_arq.is_file() or not planilha.is_file():
        raise SystemExit(f"erro: rode 'montar' antes (não achei {chave_arq} e {planilha})")
    chave = json.loads(chave_arq.read_text(encoding="utf-8"))
    if chave.get("_candidato") != pasta.name:
        raise SystemExit(f"erro: a chave é do candidato {chave.get('_candidato')!r}, não de {pasta.name!r}")
    resultado = placar(ler_planilha(planilha), chave)
    contagens_arq = saida / "contagens.json"
    contagens = json.loads(contagens_arq.read_text(encoding="utf-8")) if contagens_arq.is_file() else {}
    relatorio = args.relatorio or EVIDENCIAS / f"comparacao-{dt.date.today().isoformat()}.md"
    escrever_relatorio(relatorio, pasta.name, resultado, contagens)
    imprimir({"candidato": pasta.name, "relatorio": relatorio.name, "grupos": resultado})

    vencedores = [g for g, r in resultado.items() if r["veredito"] == "candidato melhor"]
    for g, r in resultado.items():
        print(f"{g}: {r['veredito']} (candidato {r['vitorias_candidato']} × atual {r['vitorias_atual']}, "
              f"p = {r['p_valor_sinal']})")
    if not args.promover:
        if vencedores:
            print(f"para promover {', '.join(vencedores)}: rode de novo com --promover")
        return 0

    alvo = [g for g in resultado if g in vencedores or args.forcar]
    alvo = [g for g in alvo if resultado[g]["revisados"] > 0]
    if not alvo:
        print("nada promovido: o candidato não venceu em nenhum grupo (use --forcar para promover mesmo assim)")
        return 0
    backup = DADOS / "modelos_anteriores" / dt.datetime.now().strftime("%Y-%m-%d_%H%M%S")
    for linha in promover(alvo, pasta, backup):
        print(f"  {linha}")
    if "interesse" in alvo:
        print("  interesse só liga com INTERESSE_MODELO=1 no .env do analise-service")
    print("  reinicie o analise-service para carregar o modelo novo")
    print(f"  para voltar: copie os arquivos de {backup} de volta para app/modelos/")
    return 0


def main(argv: list[str]) -> int:
    p = argparse.ArgumentParser(description="Comparação cega atual × candidato nas divergências (só contagens).")
    sub = p.add_subparsers(dest="comando", required=True)
    for nome in ("montar", "placar"):
        s = sub.add_parser(nome)
        s.add_argument("--candidato", type=Path, default=None, help="pasta do candidato (padrão: o mais recente)")
        s.add_argument("--candidatos", type=Path, default=DADOS / "candidatos")
        s.add_argument("--saida", type=Path, default=None, help="padrão dados_locais/comparacao/<candidato>/")
    m = sub.choices["montar"]
    m.add_argument("--corpus", type=Path, default=DADOS / "corpus.csv")
    m.add_argument("--holdout", type=Path, default=DADOS / "holdout_reunioes.csv")
    m.add_argument("--col-reuniao", default="ID_MEETING")
    m.add_argument("--col-texto", default="ANON_TRANSCRICAO")
    m.add_argument("--col-holdout", default="ID_MEETING")
    m.add_argument("--atual", type=Path, default=CAMINHO_MODELO_ATUAL)
    m.add_argument("--amostra", type=int, default=AMOSTRA_PADRAO, help="divergências por grupo na planilha")
    m.add_argument("--grupos", nargs="+", choices=GRUPOS, default=list(GRUPOS))
    m.add_argument("--semente", type=int, default=SEMENTE)
    pl = sub.choices["placar"]
    pl.add_argument("--relatorio", type=Path, default=None, help="padrão docs/evidencias/comparacao-<data>.md")
    pl.add_argument("--promover", action="store_true", help="copia o candidato para app/modelos/ onde ele venceu")
    pl.add_argument("--forcar", action="store_true", help="com --promover: promove mesmo sem vitória clara")
    args = p.parse_args(argv)
    return cmd_montar(args) if args.comando == "montar" else cmd_placar(args)


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
